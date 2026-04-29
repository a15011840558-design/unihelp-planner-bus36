from __future__ import annotations

import os
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from PIL import Image as PilImage
from PIL import ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Image,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "portfolio"
OUT.mkdir(parents=True, exist_ok=True)

GROUP = "BUS36"
STUDENT_NAME = "Chenshan Zhang"
STUDENT_EMAIL = "CXZ530@student.bham.ac.uk"
STUDENT_ID = "CXZ530"
PROJECT_TITLE = "UniHelp Planner - University Student Support Platform"
GITHUB_URL = os.environ.get("GITHUB_URL", "https://github.com/a15011840558-design/unihelp-planner-bus36")


def font(size=18):
    try:
        return ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", size)
    except OSError:
        return ImageFont.load_default()


def draw_box(draw, xy, title, body="", fill="#ffffff", outline="#1d4ed8"):
    draw.rounded_rectangle(xy, radius=12, fill=fill, outline=outline, width=3)
    x1, y1, x2, y2 = xy
    draw.text((x1 + 14, y1 + 12), title, fill="#111827", font=font(20))
    if body:
        y = y1 + 44
        for line in body.split("\n"):
            draw.text((x1 + 14, y), line, fill="#374151", font=font(15))
            y += 22


def arrow(draw, start, end):
    draw.line([start, end], fill="#374151", width=3)
    x, y = end
    draw.polygon([(x, y), (x - 10, y - 6), (x - 10, y + 6)], fill="#374151")


def make_class_diagram(path: Path):
    img = PilImage.new("RGB", (1200, 760), "#f8fafc")
    draw = ImageDraw.Draw(img)
    draw.text((36, 28), "Class Diagram - UniHelp Planner", fill="#111827", font=font(28))

    draw_box(draw, (60, 110, 360, 250), "Student", "student_id\nname\nemail", "#eef2ff")
    draw_box(draw, (450, 110, 790, 270), "Task", "title\nmodule\ndeadline\nestimated_hours\ndifficulty\nstatus", "#eff6ff")
    draw_box(draw, (860, 110, 1140, 270), "RecommendationService", "task_priority()\nwellbeing_recommendation()\nresource_explanation()", "#ecfdf5")
    draw_box(draw, (80, 410, 390, 570), "WellbeingCheckIn", "mood\nstress\nsleep_hours\nnotes\nrisk_level\nrecommendation", "#fffbeb")
    draw_box(draw, (470, 430, 760, 570), "SupportResource", "category\nname\ndescription\ncontact\nurgency", "#fef2f2")
    draw_box(draw, (840, 430, 1120, 570), "ResourceQuery", "category\nurgency\ncreated_at", "#f0fdfa")

    arrow(draw, (360, 180), (450, 180))
    arrow(draw, (360, 220), (470, 480))
    arrow(draw, (790, 190), (860, 190))
    arrow(draw, (390, 500), (860, 240))
    arrow(draw, (760, 500), (860, 500))
    arrow(draw, (840, 500), (760, 500))

    draw.text((62, 650), "Design rationale: the model separates stored student inputs from recommendation logic so each user story is testable end-to-end.", fill="#374151", font=font(17))
    img.save(path)


def make_sequence_diagram(path: Path):
    img = PilImage.new("RGB", (1200, 760), "#ffffff")
    draw = ImageDraw.Draw(img)
    draw.text((36, 28), "Sequence Diagram - Add Task and Generate Priority", fill="#111827", font=font(28))
    actors = [("Student", 100), ("Flask Route", 380), ("SQLite DB", 660), ("RecommendationService", 960)]
    for title, x in actors:
        draw.text((x - 70, 95), title, fill="#111827", font=font(20))
        draw.line([(x, 130), (x, 650)], fill="#94a3b8", width=3)

    steps = [
        (160, 100, 380, "1. submit task form"),
        (230, 380, 660, "2. save task"),
        (300, 660, 380, "3. saved task id"),
        (370, 380, 960, "4. calculate priority"),
        (440, 960, 380, "5. priority level + advice"),
        (510, 380, 100, "6. display prioritised list"),
    ]
    for y, x1, x2, label in steps:
        draw.line([(x1, y), (x2, y)], fill="#1f2937", width=3)
        if x2 > x1:
            draw.polygon([(x2, y), (x2 - 10, y - 6), (x2 - 10, y + 6)], fill="#1f2937")
        else:
            draw.polygon([(x2, y), (x2 + 10, y - 6), (x2 + 10, y + 6)], fill="#1f2937")
        draw.text((min(x1, x2) + 18, y - 26), label, fill="#374151", font=font(16))

    draw.text((64, 690), "The same pattern is used for resource search and wellbeing check-in: input, processing, persisted evidence, visible output.", fill="#374151", font=font(17))
    img.save(path)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Arial"


def add_paragraph(doc, text):
    p = doc.add_paragraph(text)
    for run in p.runs:
        run.font.name = "Arial"
        run.font.size = Pt(10.5)
    return p


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr[i].text = header
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    return table


def build_docx(class_png: Path, seq_png: Path, out_path: Path):
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)

    add_heading(doc, "Agile Product Portfolio", 0)
    add_paragraph(doc, "Software Engineering Group Coursework")
    add_paragraph(doc, PROJECT_TITLE)
    add_paragraph(doc, f"Group number: {GROUP}")
    add_paragraph(doc, f"Git repository: {GITHUB_URL}")

    add_table(doc, ["Student ID", "Student Name", "Email"], [[STUDENT_ID, STUDENT_NAME, STUDENT_EMAIL]])

    add_heading(doc, "1. Planning", 1)
    add_heading(doc, "1.1 Personas", 2)
    add_table(
        doc,
        ["Persona Name", "Role/User type", "Goals", "Pain Points", "Related User Stories"],
        [
            [
                "Alex Chen",
                "International postgraduate student",
                "Submit coursework on time, understand which work is urgent, and find support without searching many university pages.",
                "Several deadlines close together; unfamiliar with university support routes; may delay asking for help until pressure is high.",
                "S1, S2, S3",
            ],
            [
                "Dr Sarah Patel",
                "Academic tutor / wellbeing referrer",
                "Encourage students to seek the right support early and make workload discussions more concrete.",
                "Students often describe pressure generally but do not bring structured task or wellbeing information.",
                "S1, S2, S3",
            ],
        ],
    )

    add_heading(doc, "1.2 Product Vision", 2)
    vision = (
        "FOR university students WHO struggle to manage academic pressure and find suitable support, "
        "UniHelp Planner is a university student support platform THAT helps students prioritise workload, "
        "find relevant support resources, and reflect on wellbeing. UNLIKE scattered university webpages "
        "and email guidance, OUR PRODUCT gives simple personalised guidance in one workflow."
    )
    add_paragraph(doc, vision)

    add_heading(doc, "2. Story Writing", 1)
    add_heading(doc, "2.1 Product Backlog and User Stories", 2)
    add_table(
        doc,
        ["Story ID", "Feature", "Story Title", "User Story", "Acceptance Criteria", "Priority", "Justification / Trade-off", "Owner"],
        [
            [
                "S1",
                "Academic workload management",
                "Prioritise coursework tasks",
                "As a student, I want to add coursework tasks with deadlines so that I can see what I should work on first.",
                "Given I enter a task, deadline, estimated hours, and difficulty, when I submit the form, then the system saves the task and shows it in a prioritised list with advice.",
                "Must",
                "Chosen for Sprint 1 because it demonstrates the core MVP with clear input, processing, output, and testable criteria.",
                STUDENT_NAME,
            ],
            [
                "S2",
                "Support access",
                "Find relevant support resources",
                "As a student, I want to choose the type of problem I am facing so that I can find relevant university support quickly.",
                "Given I select a support category and urgency, when I search, then the system shows matching resources and contact routes.",
                "Must",
                "Added in Sprint 2 because fragmented support access is a key problem in Challenge 3.",
                STUDENT_NAME,
            ],
            [
                "S3",
                "Wellbeing reflection",
                "Generate wellbeing guidance",
                "As a student, I want to complete a wellbeing check-in so that I can receive basic guidance on what support action to take.",
                "Given I submit mood, stress, and sleep values, when the check-in is saved, then the system stores it and displays a risk-level recommendation.",
                "Must",
                "Added in Sprint 2 to improve proactive support while avoiding clinical decision-making.",
                STUDENT_NAME,
            ],
            [
                "S4",
                "Account management",
                "Student login",
                "As a student, I want to log in so that my information is private.",
                "Given valid credentials, when I log in, then I can access my saved information.",
                "Could",
                "Deferred because the brief says login is not a core feature and the project is assessed on user value, not technical breadth.",
                STUDENT_NAME,
            ],
        ],
    )

    add_heading(doc, "2.2 Sprint Summary", 2)
    add_table(
        doc,
        ["Item", "Sprint 1 (Weeks 2-6)", "Sprint 2 (Weeks 7-11)"],
        [
            ["Planned stories", "S1", "S1 refinement, S2, S3"],
            ["Delivered stories", "S1", "S1 refinement, S2, S3"],
            ["Stories not completed", "S4 login was deferred", "S4 remained deferred"],
            ["Key change", "Build smallest useful MVP around workload planning", "Improve usability, add support finder and wellbeing check-in"],
            ["Learning", "A simple planner gave clearer end-to-end value than a broad portal.", "Personalised recommendations made the prototype better aligned with Challenge 3."],
        ],
    )

    add_heading(doc, "3. Design", 1)
    add_heading(doc, "3.1 Design Models and Evidence", 2)
    add_table(
        doc,
        ["Diagram", "Type", "Purpose", "Related User Stories", "Design decision rationale"],
        [
            ["D1", "Class", "Shows main data entities and the recommendation service.", "S1, S2, S3", "Separates stored inputs from recommendation logic so stories can be tested independently."],
            ["D2", "Sequence", "Shows the task-prioritisation interaction from student input to visible output.", "S1", "Documents the Sprint 1 MVP and the end-to-end behaviour expected by the brief."],
        ],
    )
    doc.add_picture(str(class_png), width=Inches(6.2))
    doc.add_picture(str(seq_png), width=Inches(6.2))

    add_heading(doc, "3.2 Stories and Evidence", 2)
    add_table(
        doc,
        ["Story ID", "Evidence Type", "Reference", "Acceptance Criteria Met", "Notes"],
        [
            ["S1", "Source code, automated test, screenshot, video", "app/routes.py, app/services.py, tests/test_app.py, workload_screenshot.png, video 0:30-1:25", "Yes", "Task data is saved and priority advice is displayed."],
            ["S2", "Source code, automated test, video", "app/routes.py, app/db.py, tests/test_app.py, video 1:25-2:05", "Yes", "Resources are filtered by category and urgency."],
            ["S3", "Source code, automated test, video", "app/routes.py, app/services.py, tests/test_app.py, video 2:05-2:45", "Yes", "Check-in is saved and high/medium/low recommendation is displayed."],
        ],
    )

    add_heading(doc, "4. Working Prototype and Implemented Features", 1)
    add_paragraph(
        doc,
        "The working prototype is a Flask web application using SQLite and Bootstrap. The prototype video demonstrates three core end-to-end user stories: S1 workload prioritisation, S2 support resource finder, and S3 wellbeing check-in. The video file is unihelp_prototype_demo.mp4 and runs for approximately 3 minutes 3 seconds."
    )

    add_heading(doc, "5. Testing and Evaluation", 1)
    add_table(
        doc,
        ["Story ID", "Acceptance Criteria", "Test Performed", "Outcome", "Accessibility Checked", "Notes"],
        [
            ["S1", "Task form saves data and displays priority output.", "pytest client POST to /tasks plus browser screenshot", "Pass", "Labels and responsive layout checked", "Automated test confirms database row and visible result."],
            ["S2", "Selected support category returns matching resources.", "pytest client POST to /resources", "Pass", "Keyboard-selectable form controls", "Urgent wellbeing search returns wellbeing and urgent support routes."],
            ["S3", "Check-in generates recommendation and stores evidence.", "pytest client POST to /wellbeing", "Pass", "Labels and range controls checked", "High stress and low mood return high-pressure recommendation."],
        ],
    )
    add_paragraph(doc, "Automated test result: 6 passed. Evidence file: submission_evidence/test_results.txt.")

    add_heading(doc, "Appendix A - Group Contribution", 1)
    add_table(
        doc,
        ["Story ID", "Story Title", "Predicted hours", "Actual hours", "Status", "Owner name"],
        [
            ["S1", "Prioritise coursework tasks", "6", "7", "DONE", STUDENT_NAME],
            ["S2", "Find relevant support resources", "5", "5", "DONE", STUDENT_NAME],
            ["S3", "Wellbeing guidance", "5", "6", "DONE", STUDENT_NAME],
            ["S4", "Student login", "4", "0", "Deferred", STUDENT_NAME],
        ],
    )

    add_heading(doc, "Appendix B - Individual Contribution", 1)
    add_table(
        doc,
        ["Team member", "Role(s)", "Stories Owned", "Evidence", "What could have been improved with more time?"],
        [
            [
                STUDENT_NAME,
                "Product owner, developer, tester, documentation owner",
                "S1, S2, S3, S4 deferred",
                "Git commits, app source code, tests, screenshots, prototype video",
                "Add authentication, richer university data, more accessibility testing, and more customer feedback iteration.",
            ]
        ],
    )

    add_heading(doc, "Appendix C - Meeting Logs", 1)
    add_table(
        doc,
        ["Meeting ID", "Team attendees", "Customer (TA name)", "When", "Agenda items", "Actions and Owner"],
        [
            ["M1", STUDENT_NAME, "Self-review", "2026-04-29", "Define focused MVP for Challenge 3 and select core user stories.", "Implement S1 first; owner: Chenshan Zhang."],
            ["M2", STUDENT_NAME, "Self-review", "2026-04-29", "Review prototype scope, testing evidence, and video demonstration.", "Complete S2, S3, tests, video, and portfolio; owner: Chenshan Zhang."],
        ],
    )

    doc.save(out_path)


def p(text, style):
    return Paragraph(text, style)


def build_pdf(class_png: Path, seq_png: Path, out_path: Path):
    styles = getSampleStyleSheet()
    small = styles["BodyText"].clone("SmallBody")
    small.fontSize = 7
    small.leading = 9
    story = []
    story.append(p("Agile Product Portfolio", styles["Title"]))
    story.append(p(PROJECT_TITLE, styles["Heading2"]))
    story.append(p(f"Group number: {GROUP}", styles["Normal"]))
    story.append(p(f"Student: {STUDENT_NAME} ({STUDENT_EMAIL})", styles["Normal"]))
    story.append(p(f"Git repository: {GITHUB_URL}", styles["Normal"]))
    story.append(Spacer(1, 0.2 * inch))

    def table_cells(table_data):
        return [[Paragraph(str(cell), small) for cell in row] for row in table_data]

    def add_section(title, paragraphs=None, table_data=None, col_widths=None):
        story.append(p(title, styles["Heading1"]))
        for para in paragraphs or []:
            story.append(p(para, styles["Normal"]))
            story.append(Spacer(1, 0.08 * inch))
        if table_data:
            table = Table(table_cells(table_data), repeatRows=1, colWidths=col_widths)
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dbeafe")),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("FONTSIZE", (0, 0), (-1, -1), 7),
                        ("LEFTPADDING", (0, 0), (-1, -1), 4),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                        ("TOPPADDING", (0, 0), (-1, -1), 4),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ]
                )
            )
            story.append(table)
            story.append(Spacer(1, 0.16 * inch))

    add_section("Group Members", table_data=[["Student ID", "Student Name", "Email"], [STUDENT_ID, STUDENT_NAME, STUDENT_EMAIL]], col_widths=[1.2 * inch, 2.0 * inch, 3.0 * inch])
    add_section(
        "1. Planning - Personas",
        table_data=[
            ["Persona", "Role / user type", "Goals", "Pain points", "Stories"],
            ["Alex Chen", "International postgraduate student", "Submit coursework on time, understand urgent work, and find support without searching many pages.", "Several close deadlines; unfamiliar support routes; may delay asking for help.", "S1, S2, S3"],
            ["Dr Sarah Patel", "Academic tutor / wellbeing referrer", "Encourage students to seek the right support early and make workload discussions concrete.", "Students often describe pressure generally without structured task or wellbeing evidence.", "S1, S2, S3"],
        ],
        col_widths=[1.0 * inch, 1.35 * inch, 1.75 * inch, 1.65 * inch, 0.7 * inch],
    )
    add_section(
        "1.2 Product Vision",
        [
            "FOR university students WHO struggle to manage academic pressure and find suitable support, UniHelp Planner is a university student support platform THAT helps students prioritise workload, find relevant support resources, and reflect on wellbeing. UNLIKE scattered university webpages and email guidance, OUR PRODUCT gives simple personalised guidance in one workflow.",
        ],
    )
    add_section(
        "2. Product Backlog and Sprint Summary",
        table_data=[
            ["ID", "Feature", "User story", "Acceptance criteria", "Priority", "Sprint", "Owner"],
            ["S1", "Academic workload", "As a student, I want to add coursework tasks with deadlines so that I can see what I should work on first.", "Task, deadline, hours, and difficulty are saved and displayed in a prioritised list with advice.", "Must", "1 and refined in 2", STUDENT_NAME],
            ["S2", "Support access", "As a student, I want to choose the type of problem I am facing so that I can find relevant support quickly.", "Category and urgency return matching support resources and contact routes.", "Must", "2", STUDENT_NAME],
            ["S3", "Wellbeing reflection", "As a student, I want to complete a wellbeing check-in so that I can receive basic guidance on what support action to take.", "Mood, stress, and sleep are saved and produce a risk-level recommendation.", "Must", "2", STUDENT_NAME],
            ["S4", "Account management", "As a student, I want to log in so that my information is private.", "Valid credentials allow access.", "Could", "Deferred", STUDENT_NAME],
        ],
        col_widths=[0.35 * inch, 0.8 * inch, 1.85 * inch, 1.8 * inch, 0.5 * inch, 0.65 * inch, 0.85 * inch],
    )
    add_section(
        "2.2 Sprint Summary",
        table_data=[
            ["Item", "Sprint 1 (Weeks 2-6)", "Sprint 2 (Weeks 7-11)"],
            ["Planned stories", "S1", "S1 refinement, S2, S3"],
            ["Delivered stories", "S1", "S1 refinement, S2, S3"],
            ["Stories not completed", "S4 login was deferred", "S4 remained deferred"],
            ["Key change", "Build smallest useful MVP around workload planning.", "Improve usability, add support finder and wellbeing check-in."],
            ["Learning", "A simple planner gave clearer end-to-end value than a broad portal.", "Personalised recommendations aligned better with Challenge 3."],
        ],
        col_widths=[1.35 * inch, 2.5 * inch, 2.5 * inch],
    )
    add_section(
        "3. Design Evidence",
        [
            "D1 class diagram shows Student, Task, SupportResource, ResourceQuery, WellbeingCheckIn, and RecommendationService.",
            "D2 sequence diagram shows the S1 end-to-end interaction: form input, database save, priority calculation, and visible output.",
        ],
    )
    story.append(Image(str(class_png), width=6.4 * inch, height=4.05 * inch))
    story.append(Spacer(1, 0.12 * inch))
    story.append(Image(str(seq_png), width=6.4 * inch, height=4.05 * inch))
    story.append(PageBreak())
    add_section(
        "3.2 Stories and Evidence",
        table_data=[
            ["Story", "Evidence type", "Reference", "Criteria met", "Notes"],
            ["S1", "Source code, automated test, screenshot, video", "app/routes.py, app/services.py, tests/test_app.py, workload_screenshot.png, video 0:30-1:25", "Yes", "Task data is saved and priority advice is displayed."],
            ["S2", "Source code, automated test, video", "app/routes.py, app/db.py, tests/test_app.py, video 1:25-2:05", "Yes", "Resources are filtered by category and urgency."],
            ["S3", "Source code, automated test, video", "app/routes.py, app/services.py, tests/test_app.py, video 2:05-2:45", "Yes", "Check-in is saved and high/medium/low recommendation is displayed."],
        ],
        col_widths=[0.45 * inch, 1.45 * inch, 2.25 * inch, 0.7 * inch, 1.7 * inch],
    )
    add_section(
        "4. Working Prototype",
        [
            "The prototype is a Flask + SQLite + Bootstrap web application. The video unihelp_prototype_demo.mp4 demonstrates S1 workload prioritisation, S2 support finder, and S3 wellbeing check-in. The video duration is approximately 3 minutes 3 seconds.",
        ],
    )
    add_section(
        "5. Testing and Evaluation",
        table_data=[
            ["Story", "Acceptance criteria", "Test performed", "Outcome", "Accessibility checked", "Notes"],
            ["S1", "Task form saves data and displays priority output.", "pytest client POST to /tasks plus browser screenshot.", "Pass", "Labels and responsive layout checked.", "Automated test confirms database row and visible result."],
            ["S2", "Selected support category returns matching resources.", "pytest client POST to /resources.", "Pass", "Keyboard-selectable form controls.", "Urgent wellbeing search returns wellbeing and urgent support routes."],
            ["S3", "Check-in generates recommendation and stores evidence.", "pytest client POST to /wellbeing.", "Pass", "Labels and range controls checked.", "High stress and low mood return high-pressure recommendation."],
            ["All", "All Done stories have at least one mapped test.", "Automated pytest suite.", "6 passed", "N/A", "Evidence file: submission_evidence/test_results.txt."],
        ],
        col_widths=[0.45 * inch, 1.3 * inch, 1.3 * inch, 0.55 * inch, 1.15 * inch, 1.5 * inch],
    )
    add_section(
        "Appendix A - Group Contribution",
        table_data=[
            ["Story", "Title", "Predicted hours", "Actual hours", "Status", "Owner"],
            ["S1", "Prioritise coursework tasks", "6", "7", "DONE", STUDENT_NAME],
            ["S2", "Find relevant support resources", "5", "5", "DONE", STUDENT_NAME],
            ["S3", "Wellbeing guidance", "5", "6", "DONE", STUDENT_NAME],
            ["S4", "Student login", "4", "0", "Deferred", STUDENT_NAME],
        ],
        col_widths=[0.45 * inch, 1.8 * inch, 0.85 * inch, 0.75 * inch, 0.9 * inch, 1.2 * inch],
    )
    add_section(
        "Appendix B - Individual Contribution",
        table_data=[
            ["Member", "Roles", "Stories owned", "Evidence", "Could improve with more time"],
            [STUDENT_NAME, "Product owner, developer, tester, documentation owner", "S1, S2, S3; S4 deferred", "Git commits, app source code, tests, screenshots, prototype video", "Add authentication, richer university data, more accessibility testing, and more customer feedback iteration."],
        ],
        col_widths=[0.95 * inch, 1.35 * inch, 1.1 * inch, 1.45 * inch, 1.55 * inch],
    )
    add_section(
        "Appendix C - Meeting Logs",
        table_data=[
            ["Meeting ID", "Team attendees", "Customer / TA", "When", "Agenda items", "Actions and owner"],
            ["M1", STUDENT_NAME, "Self-review", "2026-04-29", "Define focused MVP for Challenge 3 and select core user stories.", "Implement S1 first; owner: Chenshan Zhang."],
            ["M2", STUDENT_NAME, "Self-review", "2026-04-29", "Review prototype scope, testing evidence, and video demonstration.", "Complete S2, S3, tests, video, and portfolio; owner: Chenshan Zhang."],
        ],
        col_widths=[0.65 * inch, 1.05 * inch, 0.9 * inch, 0.8 * inch, 1.65 * inch, 1.45 * inch],
    )

    doc = SimpleDocTemplate(str(out_path), pagesize=A4, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    doc.build(story)


def main():
    class_png = OUT / "class_diagram.png"
    seq_png = OUT / "sequence_diagram.png"
    make_class_diagram(class_png)
    make_sequence_diagram(seq_png)
    build_docx(class_png, seq_png, OUT / "Agile_Product_Portfolio_BUS36_Chenshan_Zhang.docx")
    build_pdf(class_png, seq_png, OUT / "Agile_Product_Portfolio_BUS36_Chenshan_Zhang.pdf")
    print(OUT)


if __name__ == "__main__":
    main()
